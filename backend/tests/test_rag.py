import io

import pytest
import pytest_asyncio
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.models import Base, Document, DocumentChunk
from app.rag.chunker import chunk_text
from app.rag.parsers import parse_docx, parse_document, parse_markdown, parse_pdf, parse_text
from app.rag.retriever import KeywordRetriever, extract_keywords


# ---------- 测试辅助：构造最小合法 PDF / DOCX 字节 ----------


def _make_pdf(*lines: str) -> bytes:
    """构造一个含文本行的最小合法 PDF（含正确 xref 表），用于解析单测。"""
    content_stream = b"".join(
        b"BT /F1 24 Tf 72 720 Td (%s) Tj ET\n" % line.encode("latin-1") for line in lines
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content_stream), content_stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i
        out += obj
        out += b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += b"%010d 00000 n \n" % off
    out += (
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF"
        % (len(objects) + 1, xref_pos)
    )
    return bytes(out)


def _make_docx(*paragraphs: str) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- 解析 ----------


def test_parse_text_collapses_blank_lines():
    assert parse_text("a\r\nb\r\n\r\n\r\nc") == "a\nb\n\nc"


def test_parse_markdown_strips_code_fence():
    md = "# 标题\n正文\n```python\nprint(1)\n```\n结尾"
    text = parse_markdown(md)
    assert "# 标题" in text
    assert "print" not in text
    assert "结尾" in text


def test_parse_pdf_extracts_text():
    data = _make_pdf("Hello PDF", "Second line")
    text = parse_pdf(data)
    assert "Hello PDF" in text
    assert "Second line" in text


def test_parse_docx_extracts_paragraphs():
    data = _make_docx("部署流程", "第一步安装依赖")
    text = parse_docx(data)
    assert "部署流程" in text
    assert "第一步安装依赖" in text


def test_parse_document_unknown_type():
    with pytest.raises(ValueError):
        parse_document("exe", b"xx")


def test_parse_document_gbk_txt():
    data = "部署流程".encode("gbk")
    assert "部署流程" in parse_document("txt", data)


def test_parse_document_md_route():
    assert "正文" in parse_document("md", "# 标题\n正文".encode("utf-8"))


# ---------- 切片 ----------


def test_chunk_text_merges_short_segments():
    text = "第一段内容。\n\n第二段内容。\n\n第三段内容。"
    chunks = chunk_text(text, max_chars=100)
    assert len(chunks) == 1
    assert "第一段内容" in chunks[0]
    assert "第三段内容" in chunks[0]


def test_chunk_text_splits_oversize():
    seg = "甲" * 300
    chunks = chunk_text(seg, max_chars=100)
    assert len(chunks) == 3
    assert all(len(c) <= 100 for c in chunks)


def test_chunk_text_heading_boundary():
    text = "导语\n# 章节一\n" + "甲" * 80 + "\n# 章节二\n" + "乙" * 80
    chunks = chunk_text(text, max_chars=100)
    assert any("章节一" in c for c in chunks)
    assert any("章节二" in c for c in chunks)
    assert not any("章节一" in c and "章节二" in c for c in chunks)


def test_chunk_text_empty():
    assert chunk_text("   \n  ") == []


# ---------- 检索 ----------


@pytest_asyncio.fixture
async def rag_db():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with factory() as db:
        yield db
    await engine.dispose()


async def _seed_docs(db, project_id="p1"):
    doc = Document(
        id="d1", tenant_id="default", project_id=project_id,
        name="部署指南.md", content_type="md",
    )
    chunks = [
        DocumentChunk(
            id=f"c{i}", tenant_id="default", project_id=project_id,
            document_id="d1", seq=i, content=text,
        )
        for i, text in enumerate(
            [
                "本文介绍如何部署项目到服务器。",
                "部署流程：先安装依赖，再启动服务。",
                "前端构建使用 npm run build。",
            ]
        )
    ]
    db.add(doc)
    db.add_all(chunks)
    await db.commit()


def test_extract_keywords_mixed():
    kws = extract_keywords("部署流程与前端 build")
    assert "build" in kws
    assert "部署" in kws
    assert "流程" in kws
    assert "前端" in kws


def test_extract_keywords_filters_stopwords():
    kws = extract_keywords("如何部署")
    assert "如何" not in kws
    assert "部署" in kws


async def test_retriever_scores_by_hits(rag_db):
    await _seed_docs(rag_db)
    retriever = KeywordRetriever(top_k=5)
    hits = await retriever.top_chunks(rag_db, "p1", "部署 流程")
    # 命中：chunk0（部署）、chunk1（部署+流程）3 个 bigram 全中；chunk2 无
    assert [h.seq for h in hits] == [1, 0]
    assert hits[0].document_name == "部署指南.md"
    assert len(hits[0].matched) == 3


async def test_retriever_limited_to_other_project(rag_db):
    await _seed_docs(rag_db)
    retriever = KeywordRetriever(top_k=5)
    assert await retriever.top_chunks(rag_db, "p2", "部署 流程") == []


async def test_retriever_no_match(rag_db):
    await _seed_docs(rag_db)
    retriever = KeywordRetriever(top_k=5)
    assert await retriever.top_chunks(rag_db, "p1", "完全不存在的词汇") == []


async def test_retriever_empty_query(rag_db):
    await _seed_docs(rag_db)
    retriever = KeywordRetriever(top_k=5)
    assert await retriever.top_chunks(rag_db, "p1", "的") == []


# ---------- API 全链路（上传/列表/删除/扫描/问答） ----------

MD_CONTENT = (
    "# 部署指南\n\n"
    "部署流程：先安装依赖，再启动服务。\n\n"
    "构建命令 npm run build。\n"
)


class _FakeAiEngine:
    """替代 AiEngine：记录 prompt 并返回固定答案。"""

    def __init__(self, answer: str = "AI 回答"):
        self.answer = answer
        self.prompts: list[str] = []

    async def chat(self, query: str, user: str = "unknown") -> str:
        self.prompts.append(query)
        return self.answer

    async def knowledge_query(self, query: str, user: str = "unknown") -> str:
        return self.answer


async def _make_project(client, headers, name="测试项目", local_path=None) -> str:
    payload: dict = {"name": name}
    if local_path:
        payload["local_path"] = str(local_path)
    r = await client.post("/api/projects", json=payload, headers=headers)
    assert r.status_code == 201
    return r.json()["id"]


async def test_upload_and_query_knowledge(client, auth_headers, monkeypatch):
    pid = await _make_project(client, auth_headers)
    r = await client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("guide.md", MD_CONTENT.encode("utf-8"), "text/markdown")},
        headers=auth_headers,
    )
    assert r.status_code == 201
    data = r.json()
    assert data["name"] == "guide.md"
    assert data["content_type"] == "md"
    assert data["status"] == "ready"

    r = await client.get(f"/api/projects/{pid}/documents", headers=auth_headers)
    assert len(r.json()) == 1

    engine = _FakeAiEngine(answer="请参考部署指南。")
    monkeypatch.setattr("app.api.knowledge.get_ai_engine", lambda: engine)
    r = await client.post(
        f"/api/projects/{pid}/knowledge", json={"query": "部署流程"}, headers=auth_headers
    )
    assert r.status_code == 200
    body = r.json()
    assert body["answer"] == "请参考部署指南。"
    assert body["sources"], "检索应命中至少一个片段"
    assert body["sources"][0]["document_name"] == "guide.md"
    assert "部署流程" in engine.prompts[0]
    assert "[guide.md #0]" in engine.prompts[0]


async def test_upload_pdf_parses(client, auth_headers):
    pid = await _make_project(client, auth_headers)
    r = await client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("doc.pdf", _make_pdf("Hello PDF"), "application/pdf")},
        headers=auth_headers,
    )
    assert r.status_code == 201
    assert r.json()["content_type"] == "pdf"
    assert r.json()["status"] == "ready"


async def test_upload_markdown_extension(client, auth_headers):
    pid = await _make_project(client, auth_headers)
    r = await client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("guide.markdown", "# 标题\n正文".encode("utf-8"), "text/markdown")},
        headers=auth_headers,
    )
    assert r.status_code == 201
    assert r.json()["content_type"] == "md"


async def test_upload_corrupt_pdf_records_error(client, auth_headers):
    pid = await _make_project(client, auth_headers)
    r = await client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("bad.pdf", b"not a real pdf", "application/pdf")},
        headers=auth_headers,
    )
    assert r.status_code == 422
    r = await client.get(f"/api/projects/{pid}/documents", headers=auth_headers)
    docs = r.json()
    assert len(docs) == 1
    assert docs[0]["status"] == "error"
    assert docs[0]["error"]


async def test_upload_unsupported_type(client, auth_headers):
    pid = await _make_project(client, auth_headers)
    r = await client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("virus.exe", b"MZ", "application/octet-stream")},
        headers=auth_headers,
    )
    assert r.status_code == 422


async def test_upload_oversize(client, auth_headers, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "rag_max_upload_mb", 0)
    pid = await _make_project(client, auth_headers)
    r = await client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("big.md", b"x", "text/markdown")},
        headers=auth_headers,
    )
    assert r.status_code == 413


async def test_upload_project_not_found(client, auth_headers):
    r = await client.post(
        "/api/projects/nope/documents",
        files={"file": ("a.md", b"x", "text/markdown")},
        headers=auth_headers,
    )
    assert r.status_code == 404


async def test_delete_document_not_found(client, auth_headers):
    pid = await _make_project(client, auth_headers)
    r = await client.delete(f"/api/projects/{pid}/documents/nope", headers=auth_headers)
    assert r.status_code == 404


async def test_knowledge_project_not_found(client, auth_headers):
    r = await client.post(
        "/api/projects/nope/knowledge", json={"query": "部署"}, headers=auth_headers
    )
    assert r.status_code == 404


async def test_scan_local_path_missing(client, auth_headers, tmp_path):
    pid = await _make_project(client, auth_headers, local_path=str(tmp_path / "nope-dir"))
    r = await client.post(f"/api/projects/{pid}/documents/scan", headers=auth_headers)
    assert r.status_code == 400


async def test_delete_document(client, auth_headers):
    pid = await _make_project(client, auth_headers)
    r = await client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("guide.md", MD_CONTENT.encode("utf-8"), "text/markdown")},
        headers=auth_headers,
    )
    doc_id = r.json()["id"]
    r = await client.delete(f"/api/projects/{pid}/documents/{doc_id}", headers=auth_headers)
    assert r.status_code == 204
    r = await client.get(f"/api/projects/{pid}/documents", headers=auth_headers)
    assert r.json() == []


async def test_scan_local_directory(client, auth_headers, tmp_path):
    (tmp_path / "README.md").write_text("部署说明：先安装依赖", encoding="utf-8")
    (tmp_path / "guide.txt").write_text("构建命令 npm run build", encoding="utf-8")
    (tmp_path / "skip.exe").write_bytes(b"MZ")
    pid = await _make_project(client, auth_headers, local_path=str(tmp_path))

    r = await client.post(f"/api/projects/{pid}/documents/scan", headers=auth_headers)
    assert r.status_code == 200
    assert r.json()["imported"] == 2
    assert r.json()["skipped"] == []

    # 二次扫描全部跳过
    r = await client.post(f"/api/projects/{pid}/documents/scan", headers=auth_headers)
    assert r.json()["imported"] == 0
    assert sorted(r.json()["skipped"]) == ["README.md", "guide.txt"]

    r = await client.get(f"/api/projects/{pid}/documents", headers=auth_headers)
    assert len(r.json()) == 2


async def test_scan_skips_unparseable_pdf(client, auth_headers, tmp_path, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "rag_scan_extensions", ["md", "txt", "markdown", "pdf"])
    (tmp_path / "ok.md").write_text("部署说明", encoding="utf-8")
    (tmp_path / "broken.pdf").write_bytes(b"not a real pdf")
    pid = await _make_project(client, auth_headers, local_path=str(tmp_path))
    r = await client.post(f"/api/projects/{pid}/documents/scan", headers=auth_headers)
    assert r.status_code == 200
    assert r.json()["imported"] == 1
    assert "broken.pdf" in r.json()["skipped"]


async def test_scan_without_local_path(client, auth_headers):
    pid = await _make_project(client, auth_headers)
    r = await client.post(f"/api/projects/{pid}/documents/scan", headers=auth_headers)
    assert r.status_code == 400


async def test_knowledge_no_match_skips_llm(client, auth_headers, monkeypatch):
    pid = await _make_project(client, auth_headers)
    r = await client.post(
        f"/api/projects/{pid}/documents",
        files={"file": ("guide.md", MD_CONTENT.encode("utf-8"), "text/markdown")},
        headers=auth_headers,
    )
    assert r.status_code == 201

    class _RaisingEngine:
        async def chat(self, *args, **kwargs):
            raise AssertionError("空检索不应调用 LLM")

    monkeypatch.setattr("app.api.knowledge.get_ai_engine", lambda: _RaisingEngine())
    r = await client.post(
        f"/api/projects/{pid}/knowledge",
        json={"query": "完全不相关的查询内容"},
        headers=auth_headers,
    )
    assert r.status_code == 200
    assert "未找到" in r.json()["answer"]
    assert r.json()["sources"] == []
